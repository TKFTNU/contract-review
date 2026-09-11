from __future__ import annotations

import html
import json
from dataclasses import asdict, dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from generate_multi_template_contract_dataset import (
    OUT_DIR as BASE_OUT_DIR,
    TEMPLATES,
    FooterCanvas,
    register_pdf_font,
    set_docx_defaults,
)


OUT_DIR = BASE_OUT_DIR / "adversarial"


@dataclass(slots=True)
class AdversarialData:
    contract_id: str
    template_key: str
    template_title: str
    template_id: str
    source_url: str
    party_a: str
    party_b: str
    city: str
    subject: str
    amount: int
    secondary_amount: int
    start_date: str
    end_date: str
    risk_type: str
    issue_codes: list[str]
    issue_summary: str


def make_record(spec: dict[str, str], serial: int) -> AdversarialData:
    key = spec["key"]
    pair_index = (serial - 1) % 2
    common = dict(
        template_key=key,
        template_title=spec["title"],
        template_id=spec["template_id"],
        source_url=spec["source_url"],
        city=["北京市", "上海市", "广州市", "深圳市", "杭州市"][serial % 5],
        party_a=["甲方测试主体", "星河数据公司", "华信咨询中心", "远航贸易公司", "安心仓储公司"][serial % 5],
        party_b=["乙方异常主体", "云图科技公司", "启明服务公司", "新域数据公司", "速达保管公司"][serial % 5],
        risk_type="adversarial",
    )
    if key == "housing_rent":
        if pair_index == 0:
            return AdversarialData(contract_id=f"ADV-{spec['template_id'].replace('-', '')}-{serial:04d}", subject="无权属证明地下室", amount=0, secondary_amount=999999, start_date="2026年12月1日", end_date="2026年11月30日", issue_codes=["date_reversed", "zero_rent", "unverified_property", "unilateral_entry"], issue_summary="租赁期限倒置、租金为零但押金极高、出租人无权属证明且可单方进入房屋", **common)
        return AdversarialData(contract_id=f"ADV-{spec['template_id'].replace('-', '')}-{serial:04d}", subject="市中心整层办公楼", amount=99999999, secondary_amount=0, start_date="2026年10月1日", end_date="2048年10月1日", issue_codes=["term_over_20y", "extreme_rent", "zero_deposit"], issue_summary="租赁期限超过二十年、月租金极端异常且约定零押金", **common)
    if key == "entrust":
        if pair_index == 0:
            return AdversarialData(contract_id=f"ADV-{spec['template_id'].replace('-', '')}-{serial:04d}", subject="代签全部合同并处分甲方资产", amount=-5000, secondary_amount=0, start_date="2026年12月1日", end_date="2026年11月1日", issue_codes=["negative_fee", "date_reversed", "unlimited_authority"], issue_summary="委托报酬为负数、期限倒置，并授予受托人无限处分和代签权限", **common)
        return AdversarialData(contract_id=f"ADV-{spec['template_id'].replace('-', '')}-{serial:04d}", subject="虚构项目的审计报告", amount=1, secondary_amount=1000000, start_date="2026年10月1日", end_date="2026年10月2日", issue_codes=["fee_mismatch", "impossible_deadline", "no_report"], issue_summary="一天内完成复杂审计、报酬仅一元但费用一百万元且无需提交成果", **common)
    if key == "intermediary":
        if pair_index == 0:
            return AdversarialData(contract_id=f"ADV-{spec['template_id'].replace('-', '')}-{serial:04d}", subject="保证百分之百成交的融资项目", amount=100000, secondary_amount=100000, start_date="2026年12月1日", end_date="2026年11月30日", issue_codes=["success_guarantee", "commission_100pct", "date_reversed"], issue_summary="中介人保证百分之百成交、收取标的额百分之百报酬且服务期限倒置", **common)
        return AdversarialData(contract_id=f"ADV-{spec['template_id'].replace('-', '')}-{serial:04d}", subject="不存在的海外工程项目", amount=500000, secondary_amount=-100, start_date="2026年10月1日", end_date="2026年10月31日", issue_codes=["negative_commission", "nonexistent_subject", "cash_only"], issue_summary="中介标的无法核验、报酬为负数且要求现金交付不留凭证", **common)
    if key == "data_provide":
        if pair_index == 0:
            return AdversarialData(contract_id=f"ADV-{spec['template_id'].replace('-', '')}-{serial:04d}", subject="含身份证号和未成年人轨迹的个人信息全集", amount=-1, secondary_amount=999999999, start_date="2026年12月1日", end_date="2099年12月1日", issue_codes=["negative_fee", "personal_data_no_consent", "perpetual_license", "public_disclosure"], issue_summary="未经同意提供敏感个人信息，负价交易、永久许可并允许公开发布", **common)
        return AdversarialData(contract_id=f"ADV-{spec['template_id'].replace('-', '')}-{serial:04d}", subject="空白数据集", amount=100, secondary_amount=0, start_date="2026年12月1日", end_date="2026年11月1日", issue_codes=["zero_records", "date_reversed", "insecure_delivery"], issue_summary="记录数为零、期限倒置，并通过未加密U盘传输数据", **common)
    if pair_index == 0:
        return AdversarialData(contract_id=f"ADV-{spec['template_id'].replace('-', '')}-{serial:04d}", subject="现金及贵重珠宝一批", amount=0, secondary_amount=0, start_date="2026年10月1日", end_date="2026年10月2日", issue_codes=["zero_fee", "zero_quantity", "liability_exemption"], issue_summary="保管费和数量均为零，却声称保管巨额贵重物品并完全免除保管人责任", **common)
    return AdversarialData(contract_id=f"ADV-{spec['template_id'].replace('-', '')}-{serial:04d}", subject="一百万件精密仪器", amount=500, secondary_amount=1000000, start_date="2026年12月1日", end_date="2026年11月1日", issue_codes=["date_reversed", "quantity_value_mismatch", "retrieval_before_storage"], issue_summary="取回日期早于入库日期、数量与费用明显不匹配且无验收记录", **common)


def clauses(data: AdversarialData) -> list[tuple[str, str]]:
    A, B = data.party_a, data.party_b
    k = data.template_key
    if k == "housing_rent":
        return [("heading", "第一条 房屋与权属"), ("body", f"甲方（出租人）：{A}；乙方（承租人）：{B}。租赁标的为{data.subject}，甲方声明没有任何权属证明。"), ("heading", "第二条 租赁期限"), ("body", f"租赁期限自{data.start_date}起至{data.end_date}止，乙方同意期限倒置仍然有效。"), ("heading", "第三条 租金与押金"), ("body", f"月租金为人民币{data.amount}元，押金为人民币{data.secondary_amount}元。"), ("heading", "第四条 特别权利"), ("body", "甲方可不经通知随时进入房屋、移动乙方物品并没收全部财物，乙方不得提出异议。"), ("heading", "第五条 争议解决"), ("body", f"乙方不得解除合同或要求退款，争议由{data.city}人民法院依法处理。"), ("body", f"甲方（签章）：{A}    乙方（签章）：{B}")]
    if k == "entrust":
        return [("heading", "第一条 委托事项与权限"), ("body", f"甲方（委托人）：{A}授权乙方（受托人）：{B}代签全部合同并处分甲方任何资产，委托事项为{data.subject}。"), ("heading", "第二条 委托期限"), ("body", f"期限自{data.start_date}起至{data.end_date}止，乙方在期限倒置时仍须完成全部事项。"), ("heading", "第三条 报酬和费用"), ("body", f"委托报酬为人民币{data.amount}元，必要费用为人民币{data.secondary_amount}元，乙方无需提供任何凭证。"), ("heading", "第四条 成果交付"), ("body", "乙方无需提交报告、成果或进度记录，甲方不得查阅办理材料。"), ("heading", "第五条 责任免除"), ("body", "无论乙方是否故意或重大过失，甲方均放弃全部追责、退款和解除权。"), ("body", f"甲方（签章）：{A}    乙方（签章）：{B}")]
    if k == "intermediary":
        return [("heading", "第一条 中介标的"), ("body", f"甲方（委托人）：{A}委托乙方（中介人）：{B}撮合{data.subject}，但乙方无需证明项目真实存在。"), ("heading", "第二条 成交保证"), ("body", "乙方保证百分之百成交并保证交易绝对盈利，即使相对方不存在也视为中介成功。"), ("heading", "第三条 服务期限"), ("body", f"服务期限自{data.start_date}起至{data.end_date}止。"), ("heading", "第四条 报酬"), ("body", f"主交易金额为人民币{data.amount}元，中介报酬为人民币{data.secondary_amount}元；报酬可为负数但甲方仍须现金支付。"), ("heading", "第五条 凭证与责任"), ("body", "乙方无需提供发票、交易记录或服务凭证，且不承担任何因虚假项目造成的责任。"), ("body", f"甲方（签章）：{A}    乙方（签章）：{B}")]
    if k == "data_provide":
        return [("heading", "第一条 数据标的"), ("body", f"甲方（接收方）：{A}，乙方（提供方）：{B}。数据标的为{data.subject}，预计记录数{data.secondary_amount}条。"), ("heading", "第二条 提供和许可期限"), ("body", f"数据于{data.start_date}通过未加密U盘交付，许可期限至{data.end_date}，乙方同意甲方永久转售并公开发布。"), ("heading", "第三条 费用"), ("body", f"数据服务费为人民币{data.amount}元，金额为负数时由乙方倒贴给甲方。"), ("heading", "第四条 合规责任"), ("body", "无需取得数据主体同意，无需脱敏，不受个人信息保护、数据安全或网络安全要求约束。"), ("heading", "第五条 违约"), ("body", "即使发生泄露或滥用，乙方也不得暂停服务、追究责任或要求删除数据。"), ("body", f"甲方（签章）：{A}    乙方（签章）：{B}")]
    return [("heading", "第一条 保管物"), ("body", f"甲方（寄存人）：{A}将{data.subject}交由乙方（保管人）：{B}保管，数量为{data.secondary_amount}件。"), ("heading", "第二条 保管期限"), ("body", f"保管期限自{data.start_date}起至{data.end_date}止，乙方应在入库前完成返还。"), ("heading", "第三条 保管费用"), ("body", f"保管费为人民币{data.amount}元，双方不进行数量、外观和价值验收。"), ("heading", "第四条 责任"), ("body", "无论因故意、重大过失、丢失或毁损，乙方均不承担任何赔偿责任。"), ("heading", "第五条 取回"), ("body", "甲方可在任何时间要求取回，乙方无需核验身份、保管单或数量。"), ("body", f"甲方（签章）：{A}    乙方（签章）：{B}")]


def add_summary_docx(doc: Document, data: AdversarialData) -> None:
    rows = [("合同编号", data.contract_id, "模板编号", data.template_id), ("模板类型", data.template_title, "数据类型", "异常测试样本"), ("甲方", data.party_a, "乙方", data.party_b), ("标的", data.subject, "金额", f"{data.amount}元"), ("期限", f"{data.start_date}—{data.end_date}", "异常代码", ", ".join(data.issue_codes))]
    table = doc.add_table(rows=len(rows), cols=4); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, vals in zip(table.rows, rows):
        for cell, value in zip(row.cells, vals):
            cell.text = value; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs: p.paragraph_format.space_after = Pt(0)
    for row in table.rows: row.cells[0].width = Cm(2.4); row.cells[2].width = Cm(2.4)


def create_docx(data: AdversarialData, path: Path) -> None:
    doc = Document(); set_docx_defaults(doc)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(f"异常测试 {data.template_title}"); r.bold = True; r.font.size = Pt(16)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f"基于 {data.template_id} 结构生成，仅用于模型压力测试").italic = True
    add_summary_docx(doc, data)
    note = doc.add_paragraph("警示：本合同故意包含明显不合理、违法风险或自相矛盾的条款，不具有法律效力，不得用于真实交易。"); note.runs[0].font.size = Pt(8.5)
    for kind, text in clauses(data):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8 if kind == "heading" else 4)
        if kind == "heading":
            p.style = doc.styles["Heading 2"]; r = p.add_run(text); r.font.name = "黑体"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        else: p.add_run(text)
    footer = doc.sections[0].footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; footer.add_run(f"{data.contract_id} | Adversarial test | No legal effect").font.size = Pt(8)
    doc.save(path)


def create_pdf(data: AdversarialData, path: Path) -> None:
    font = register_pdf_font(); styles = getSampleStyleSheet()
    title = ParagraphStyle("AdvTitle", parent=styles["Title"], fontName=font, fontSize=16, leading=22, alignment=TA_CENTER, spaceAfter=8)
    sub = ParagraphStyle("AdvSub", parent=styles["Normal"], fontName=font, fontSize=8.5, leading=12, alignment=TA_CENTER, textColor=colors.HexColor("#666666"), spaceAfter=10)
    head = ParagraphStyle("AdvHead", parent=styles["Heading2"], fontName=font, fontSize=11.5, leading=16, alignment=TA_LEFT, spaceBefore=8, spaceAfter=5)
    body = ParagraphStyle("AdvBody", parent=styles["BodyText"], fontName=font, fontSize=9.2, leading=15, alignment=TA_LEFT, spaceAfter=4)
    note = ParagraphStyle("AdvNote", parent=body, fontSize=7.8, leading=11, textColor=colors.HexColor("#8A1C1C"))
    cell = ParagraphStyle("AdvCell", parent=body, fontName=font, fontSize=7.7, leading=9.5, spaceAfter=0)
    story = [Paragraph(html.escape(f"异常测试 {data.template_title}"), title), Paragraph(html.escape(f"基于 {data.template_id} 结构生成，仅用于模型压力测试"), sub)]
    rows = [["合同编号", data.contract_id, "模板编号", data.template_id], ["模板类型", data.template_title, "数据类型", "异常测试样本"], ["甲方", data.party_a, "乙方", data.party_b], ["标的", data.subject, "金额", f"{data.amount}元"], ["期限", f"{data.start_date}—{data.end_date}", "异常代码", ", ".join(data.issue_codes)]]
    tdata = [[Paragraph(html.escape(str(v)), cell) for v in row] for row in rows]
    table = Table(tdata, colWidths=[2.2 * cm, 5.1 * cm, 2.2 * cm, 5.1 * cm]); table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9D9D9")), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#FBEAEA")), ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#FBEAEA")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story += [table, Spacer(1, 8), Paragraph("警示：本合同故意包含明显不合理、违法风险或自相矛盾的条款，不具有法律效力，不得用于真实交易。", note), Spacer(1, 6)]
    for kind, text in clauses(data): story.append(Paragraph(html.escape(text), head if kind == "heading" else body))
    def factory(*args, **kwargs): return FooterCanvas(*args, contract_id=data.contract_id, font_name=font, **kwargs)
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=1.8 * cm, leftMargin=1.8 * cm, topMargin=1.8 * cm, bottomMargin=1.7 * cm, title=f"{data.contract_id} {data.template_title}", author="Adversarial synthetic contract dataset generator")
    doc.build(story, canvasmaker=factory)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True); records = []
    serial = 1
    for spec in TEMPLATES:
        for _ in range(2):
            data = make_record(spec, serial); docx_path = OUT_DIR / f"{data.contract_id}.docx"; pdf_path = OUT_DIR / f"{data.contract_id}.pdf"
            create_docx(data, docx_path); create_pdf(data, pdf_path)
            records.append({**asdict(data), "docx": docx_path.name, "pdf": pdf_path.name}); serial += 1
    manifest = {"dataset_name": "samr_multi_template_adversarial_contracts", "count": len(records), "template_count": len(TEMPLATES), "source_note": "异常测试合同依据国家市场监督管理总局合同示范文本结构合成，故意加入明显不合理条款；所有主体和数值均为虚构。", "records": records}
    (OUT_DIR / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT_DIR / "README.md").write_text("# 异常合同测试集\n\n本目录包含10份DOCX和10份PDF异常合同，每种官方示范文本各2份。样本故意包含金额异常、期限倒置、无限授权、隐私违规、责任完全免除等问题，用于测试条款切分、矛盾检测和LLM复核能力。\n\n所有文件均为合成数据，不具有法律效力，不得用于真实交易。详见 manifest.json。\n", encoding="utf-8")
    print(json.dumps({"output_dir": str(OUT_DIR), "count": len(records)}, ensure_ascii=False))


if __name__ == "__main__": main()
